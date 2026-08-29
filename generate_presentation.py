from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

NAVY = RGBColor(0, 64, 133)
DARK = RGBColor(30, 41, 59)
LIGHT_BG = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
ACCENT_BLUE = RGBColor(37, 99, 235)
BORDER_COLOR = RGBColor(226, 232, 240)

def add_header(slide, title_text, category="ABV-IIITM GWALIOR | TECHNICAL PROJECT PROPOSAL"):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()

    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_c = cat_box.text_frame
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.size = Pt(10.5)
    p_c.font.bold = True
    p_c.font.color.rgb = ACCENT_BLUE

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = NAVY

def add_card(slide, left, top, width, height, title, items):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1)

    tx = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
    tf = tx.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(13.5)
    p0.font.bold = True
    p0.font.color.rgb = NAVY
    p0.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK
        p.space_after = Pt(5)

# SLIDE 1: Title
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = NAVY
bg1.line.fill.background()

t_box1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
tf1 = t_box1.text_frame
p1 = tf1.paragraphs[0]
p1.text = "ABV-IIITM Campus AI Assistant"
p1.font.size = Pt(38)
p1.font.bold = True
p1.font.color.rgb = WHITE
p1.space_after = Pt(10)

p2 = tf1.add_paragraph()
p2.text = "Zero-Hallucination Retrieval-Augmented Generation (RAG) System for Institute Inquiries"
p2.font.size = Pt(17)
p2.font.color.rgb = RGBColor(191, 219, 254)
p2.space_after = Pt(24)

p3 = tf1.add_paragraph()
p3.text = "Presented By: Student Technical Lead | Academic & Administrative Review 2026"
p3.font.size = Pt(13)
p3.font.color.rgb = RGBColor(148, 163, 184)

# SLIDE 2: Problem Statement
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Institutional Challenges & Motivation")
add_card(s2, 0.8, 1.8, 3.6, 5.0, "Administrative Overhead", [
    "High volume of repetitive queries during admission cycles.",
    "Academic staff spends hours clarifying identical fee & grading rules.",
    "Staff diverted from high-priority research and institutional operations."
])
add_card(s2, 4.8, 1.8, 3.6, 5.0, "Fragmented Information", [
    "Critical rules dispersed across 680+ pages of disjointed PDFs.",
    "Students miss essential clauses for attendance & branch change.",
    "Search engines fail to pinpoint specific regulatory sub-clauses."
])
add_card(s2, 8.8, 1.8, 3.6, 5.0, "Stakeholder Experience", [
    "Prospective parents struggle to find exact 2026 fee breakdowns.",
    "Lack of 24/7 self-service guidance for international/DASA applicants.",
    "Need for an instant, verifiable digital helpdesk on the portal."
])

# SLIDE 3: System Architecture
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "System Architecture: RAG Knowledge Engine")
add_card(s3, 0.8, 1.8, 5.6, 5.0, "Local Knowledge Ingestion Pipeline", [
    "687 official institute PDF pages ingested and parsed.",
    "Recursive character chunking (600 tokens with 100 overlap).",
    "2,822 semantic vectors generated via sentence-transformers.",
    "Persistent local vector storage using ChromaDB.",
    "Zero institutional documents sent to external third-party databases."
])
add_card(s3, 6.8, 1.8, 5.6, 5.0, "Retrieval & Response Pipeline", [
    "User submits query via responsive web widget.",
    "Vector similarity search extracts top 5 verified context chunks.",
    "High-speed LLM synthesizes response strictly from retrieved data.",
    "Zero-hallucination guardrail: politely routes unindexed queries to official emails.",
    "Sub-second response delivery via asynchronous FastAPI backend."
])

# SLIDE 4: Coverage Scope
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Institutional Knowledge Domains Indexed")
add_card(s4, 0.8, 1.8, 3.6, 2.4, "Admissions & Intake", [
    "Annual intake across B.Tech, M.Tech, MBA, PhD.",
    "Non-GATE M.Tech exam schedules & syllabi.",
    "MBA eligibility (CAT, MAT, CMAT, CUET)."
])
add_card(s4, 4.8, 1.8, 3.6, 2.4, "Academic Ordinances", [
    "Mandatory 75% minimum attendance rule.",
    "10-point letter grading & SGPA/CGPA math.",
    "UG branch change criteria (CGPA >= 8.00)."
])
add_card(s4, 8.8, 1.8, 3.6, 2.4, "Fee Schedules", [
    "2022 to 2026 batch tuition & hostel tables.",
    "Itemized breakdown (Rs. 1,68,750 for 2026 B.Tech).",
    "SC/ST, PwD, and DASA category fee waivers."
])
add_card(s4, 0.8, 4.5, 3.6, 2.3, "NEP 2020 Multi-Exits", [
    "Certificate in Engg. Sciences (Year 1).",
    "Diploma in CSE / IT / EEE (Year 2).",
    "B.Sc. Degree (Year 3) & B.Tech / PG."
])
add_card(s4, 4.8, 4.5, 3.6, 2.3, "Hostel & Student Life", [
    "10:00 PM compulsory night attendance.",
    "200W electrical appliance power cap.",
    "Sports complex, library, and health center."
])
add_card(s4, 8.8, 4.5, 3.6, 2.3, "Placements & Research", [
    "67.5 LPA highest CTC & average packages.",
    "80+ visiting recruiters (Google, D.E. Shaw).",
    "Final-year 90-day internship regulations."
])

# SLIDE 5: Validation
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Accuracy & Ground Truth Validation")
add_card(s5, 0.8, 1.8, 5.6, 5.0, "Test 1: 2026 B.Tech Fee Inquiry", [
    "Query: 'What is the total fee payable for B.Tech in 2026?'",
    "Extracted Answer: Rs. 1,68,750 total admission payment.",
    "Breakdown: Rs. 1,05,000 Institute fee + Rs. 35,000 Hostel fee + Rs. 23,750 One-time fee + Rs. 5,000 Mess deposit.",
    "Validation: 100% verified against official Annexure-A fee circular."
])
add_card(s5, 6.8, 1.8, 5.6, 5.0, "Test 2: Non-GATE M.Tech Admissions", [
    "Query: 'Where is the Non-GATE written test conducted?'",
    "Extracted Answer: Room 106, Lecture Theatre I (LT-I) at 10:00 AM on August 8, 2026.",
    "Validation: 100% verified against Advertisement No. 623.",
    "Test 3: Attendance Rule -> Accurately extracted 75% cutoff."
])

# SLIDE 6: Value & Rollout
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Implementation Plan & Institutional Value")
add_card(s6, 0.8, 1.8, 5.6, 5.0, "Deployment & Infrastructure", [
    "Embeds in 1 minute using a single <script> tag on iiitm.ac.in.",
    "Runs on existing campus Linux hardware / VM (Rs. 0/month cost).",
    "Automated weekly script re-indexes new circulars & notifications.",
    "Zero modifications required to existing CMS backend."
])
add_card(s6, 6.8, 1.8, 5.6, 5.0, "Institutional Value & Next Steps", [
    "Reduces routine administrative inquiry tickets by up to 70%.",
    "Provides 24/7 instant guidance for students and parents worldwide.",
    "Guarantees accurate, policy-grounded communication at all times.",
    "Next Step: Approve staging deployment on assistant.iiitm.ac.in."
])

prs.save("ABV_IIITM_Campus_AI_Assistant_Presentation.pptx")
print("SUCCESS: Slide deck saved -> ABV_IIITM_Campus_AI_Assistant_Presentation.pptx")
